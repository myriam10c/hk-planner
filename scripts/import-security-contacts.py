"""
Import building security contacts from the Hostaway Listing Maps xlsx and emit
security-contacts.json at the repo root.

Re-download source:
  curl -sL "https://drive.google.com/uc?export=download&id=1HLK8q1g6MAAs8Pvzc48vwYrYOzuE6K0i" \
       -o /tmp/security_contacts.xlsx

Usage:
  python3 scripts/import-security-contacts.py [path-to-xlsx]
  (default: /tmp/security_contacts.xlsx)

Output written to: security-contacts.json (repo root, next to this script's parent dir)
"""

import sys
import json
import re
import warnings
from pathlib import Path
from datetime import date
from collections import defaultdict

# Suppress openpyxl "Unknown extension" warnings
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
import openpyxl  # noqa: E402

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
SHEET_NAME = "Hostaway Listing Maps (3)"
DEFAULT_XLSX = "/tmp/security_contacts.xlsx"
REPO_ROOT = Path(__file__).parent.parent
OUTPUT_PATH = REPO_ROOT / "security-contacts.json"

# Column indexes (0-based relative to the 15 columns A-O)
COL_LISTING_MAP_ID = 0
COL_APT_BUILDING = 1
COL_UNIT = 2
COL_BUILDING_FROM_NAME = 3
COL_UNIT_NORM = 4
COL_BUILDING_NORM = 5
COL_EMAIL = 6
COL_EMAIL2 = 7
COL_EMAIL3 = 8
COL_NUMBER = 9
COL_PORTAL = 10
COL_PORTAL_ID = 11
COL_PORTAL_PASSWORD = 12
COL_DTCM = 13
COL_DET = 14

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def clean(v):
    """Return stripped string or None."""
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def normalise_email(s):
    """
    Return cleaned email string (lowercase domain, original-case local part).
    Does NOT validate beyond '@' + no-spaces check.
    """
    s = s.strip()
    at = s.index("@")
    local = s[:at]
    domain = s[at + 1:].lower()
    return f"{local}@{domain}"


def parse_phone(raw):
    """
    Return (phone_display, phone_digits) or (None, None) if invalid.
    - phone_display: original stripped value (the first number if comma/colon-separated)
    - phone_digits: digits only, normalised to start with country code (no leading 00)

    Handles:
      - Multiple numbers separated by , / : -> take first one only
      - Numeric cell values (float like 971509179314.0)
      - '#ERROR' -> drop
      - Leading '00' -> drop
      - 10-digit starting with '0' -> assume UAE, replace 0 -> 971
      - 9-digit starting with '5' -> prefix 971
      - Fewer than 8 digits -> drop
    """
    if raw is None:
        return None, None
    s = str(raw).strip()
    if not s:
        return None, None
    if "#ERROR" in s.upper():
        return None, None

    # If the cell is a float (numeric), convert to int string first
    # e.g. 971509179314.0 -> '971509179314'
    try:
        f = float(s)
        if f == int(f):
            s = str(int(f))
    except ValueError:
        pass

    # Multiple numbers -> take first segment
    first = re.split(r"[,;/]|(?<!\+):", s)[0].strip()

    digits = re.sub(r"[^\d]", "", first)

    if len(digits) < 8:
        return None, None

    # Leading 00 -> strip
    if digits.startswith("00"):
        digits = digits[2:]

    # 10-digit starting with 0 -> assume UAE mobile
    if digits.startswith("0") and len(digits) == 10:
        digits = "971" + digits[1:]

    # 9-digit starting with 5 -> UAE mobile missing country code
    if digits.startswith("5") and len(digits) == 9:
        digits = "971" + digits

    # Still fewer than 8 -> discard
    if len(digits) < 8:
        return None, None

    return first, digits


def classify_email_col(v):
    """
    Classify a value from email columns (6,7,8).
    Returns: ("email", normalised_email) | ("portal", url) | ("note", text) | None
    """
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if "@" in s and " " not in s:
        return ("email", normalise_email(s))
    if s.lower().startswith("http"):
        return ("portal", s)
    return ("note", s)


def normalize_bkey(raw_key):
    """
    Normalize a raw building key:
      1. Strip trailing punctuation, slashes, and whitespace.
      2. Collapse multiple internal spaces to one.
    Then apply BUILDING_ALIASES to unify known variant spellings.
    """
    k = raw_key.strip(" /\\.,;:")
    k = re.sub(r" {2,}", " ", k)
    return BUILDING_ALIASES.get(k, k)


# Known variant keys -> canonical key.
# Add new entries here whenever a new spelling variant appears in the source data.
# Keys are already lowercased + stripped (the normalized form before alias lookup).
BUILDING_ALIASES = {
    # "15th northside" variants
    "15northside":     "15th northside",
    "15th nortside":   "15th northside",
    # Azizi Riviera typo
    "azizi riveria 39": "azizi riviera 39",
    # "marwa height/" trailing slash is handled by normalize_bkey stripping,
    # so it becomes "marwa height" and merges naturally -- no alias needed.
}


def natural_sort_key(s):
    """Sort key for natural (human) ordering of unit codes like 2510B, 807, 1A."""
    parts = re.split(r"(\d+)", str(s))
    return [int(p) if p.isdigit() else p.lower() for p in parts]


def group_signature(emails, phone_digits, portal, portal_id, portal_password, notes):
    """Hashable signature for grouping rows with identical contact details."""
    return (
        tuple(sorted(emails)),
        phone_digits,
        portal,
        portal_id,
        portal_password,
        tuple(sorted(notes)),
    )


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build_contacts(xlsx_path):
    """Parse the xlsx and return (output_dict, stats_dict). No file written.

    Also used by the VPS job scheduled/hk_security_contacts.py (vendored copy).
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME]

    rows_read = 0
    rows_dropped = []  # list of (row_number, reason)

    # building_key -> list of parsed row dicts
    building_rows = defaultdict(list)
    # building_key -> {building_from_name -> count}  (for display name selection)
    building_name_counts = defaultdict(lambda: defaultdict(int))

    for r in range(2, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, 16)]

        # Skip fully empty rows
        if all(v is None or str(v).strip() == "" for v in cells):
            continue
        rows_read += 1

        # ---- Unit ----------------------------------------------------------
        unit_raw = clean(cells[COL_UNIT_NORM]) or clean(cells[COL_UNIT])
        if unit_raw is None:
            rows_dropped.append((r, "no unit"))
            continue
        # Numeric values from openpyxl come as float
        try:
            f = float(unit_raw)
            if f == int(f):
                unit_raw = str(int(f))
        except (ValueError, TypeError):
            pass
        unit = unit_raw.upper().strip()

        # ---- Building key --------------------------------------------------
        bn_raw = clean(cells[COL_BUILDING_NORM])
        bfn_raw = clean(cells[COL_BUILDING_FROM_NAME])

        if bn_raw:
            bkey = normalize_bkey(bn_raw.lower())
        elif bfn_raw:
            bkey = normalize_bkey(bfn_raw.lower())
        else:
            bkey = "unknown"

        # ---- Email columns (6,7,8) -----------------------------------------
        emails = []
        notes = []
        portal_from_email = None

        for ci in [COL_EMAIL, COL_EMAIL2, COL_EMAIL3]:
            result = classify_email_col(cells[ci])
            if result is None:
                continue
            kind, val = result
            if kind == "email":
                if val not in emails:
                    emails.append(val)
            elif kind == "portal":
                if portal_from_email is None:
                    portal_from_email = val
            elif kind == "note":
                if val not in notes:
                    notes.append(val)

        # ---- Phone ---------------------------------------------------------
        phone_display, phone_digits = parse_phone(cells[COL_NUMBER])

        # ---- Portal (col 10) -----------------------------------------------
        portal_raw = clean(cells[COL_PORTAL])
        portal_url = None
        if portal_raw:
            if portal_raw.lower().startswith("http"):
                portal_url = portal_raw
            elif "@" in portal_raw and " " not in portal_raw:
                # It's an email masquerading as portal -> treat as note
                notes.append(portal_raw)
            else:
                # Non-http portal label -> note (e.g. "Adda Application")
                notes.append(portal_raw)

        # Fall back to email-col portal if portal col empty
        if portal_url is None and portal_from_email is not None:
            portal_url = portal_from_email

        # ---- Portal ID / Password ------------------------------------------
        portal_id = clean(cells[COL_PORTAL_ID])
        portal_password = clean(cells[COL_PORTAL_PASSWORD])

        # ---- DTCM links ----------------------------------------------------
        dtcm_raw = clean(cells[COL_DTCM])
        dtcm = dtcm_raw if dtcm_raw and dtcm_raw.lower().startswith("http") else None

        det_raw = clean(cells[COL_DET])
        det = det_raw if det_raw and det_raw.lower().startswith("http") else None

        # ---- Validate: must have at least one contact signal ---------------
        has_contact = bool(emails or phone_digits or portal_url or notes)
        if not has_contact:
            rows_dropped.append((r, "no contact info"))
            continue

        # ---- Deduplicate notes ---------------------------------------------
        seen_notes = []
        for n in notes:
            if n not in seen_notes:
                seen_notes.append(n)
        notes = seen_notes

        parsed = {
            "row": r,
            "unit": unit,
            "bkey": bkey,
            "bfn": bfn_raw or "",
            "emails": emails,
            "phone": phone_display,
            "phoneDigits": phone_digits,
            "portal": portal_url,
            "portalId": portal_id,
            "portalPassword": portal_password,
            "notes": notes,
            "dtcm": dtcm,
            "det": det,
        }
        building_rows[bkey].append(parsed)
        building_name_counts[bkey][bfn_raw or bkey] += 1

    # -------------------------------------------------------------------------
    # Merge plural variants
    # (e.g. "bloom height" and "bloom heights" -> keep the one with more rows)
    # -------------------------------------------------------------------------
    all_keys = list(building_rows.keys())
    merged_into = {}  # source_key -> canonical_key

    for key in sorted(all_keys):
        singular = key
        plural = key + "s"
        if plural in building_rows and singular in building_rows:
            count_s = len(building_rows[singular])
            count_p = len(building_rows[plural])
            # Canonical = the one with more rows (ties: keep singular)
            if count_p > count_s:
                canonical, minor = plural, singular
            else:
                canonical, minor = singular, plural
            # Move minor rows into canonical
            building_rows[canonical].extend(building_rows.pop(minor))
            for bfn, cnt in building_name_counts[minor].items():
                building_name_counts[canonical][bfn] += cnt
            del building_name_counts[minor]
            merged_into[minor] = canonical

    # -------------------------------------------------------------------------
    # Build output
    # -------------------------------------------------------------------------
    buildings_out = []
    units_kept = 0

    for bkey in sorted(building_rows.keys()):
        rows = building_rows[bkey]

        # Display name: most frequent building_from_name among merged rows.
        # Tie-break: prefer the shorter name (e.g. "555 Parkview" over "555 Parkviews").
        name_counts = building_name_counts[bkey]
        display_name = max(name_counts, key=lambda k: (name_counts[k], -len(k)))
        if not display_name or display_name == bkey:
            # Fallback: title-case the key
            display_name = bkey.title()

        # Group rows by signature
        groups_map = defaultdict(list)  # signature -> list of row dicts
        for row_data in rows:
            sig = group_signature(
                row_data["emails"],
                row_data["phoneDigits"],
                row_data["portal"],
                row_data["portalId"],
                row_data["portalPassword"],
                row_data["notes"],
            )
            groups_map[sig].append(row_data)

        groups_out = []
        for sig, group_rows in groups_map.items():
            # Take contact info from first row of the group (all identical per sig)
            rep = group_rows[0]

            # ------------------------------------------------------------------
            # Deduplicate rows by canonical unit code (I-1).
            # Two unit codes are "twins" if they are identical after removing all
            # non-alphanumeric characters, e.g. "G03" and "G-03" both reduce to
            # "G03".  We keep the shorter/canonical form (G03) and merge links:
            # if the first occurrence lacks a dtcm/det link that a later duplicate
            # has, we keep the non-null link.
            # ------------------------------------------------------------------
            canonical_map = {}  # canonical_unit -> row_dict (merged)
            canonical_order = []  # insertion order for stable sorting

            for rd in sorted(group_rows, key=lambda x: natural_sort_key(x["unit"])):
                raw_unit = rd["unit"]
                canon_unit = re.sub(r"[^A-Z0-9]", "", raw_unit)  # keep only alnum
                # Prefer the shorter form (e.g. "G03" over "G-03")
                if canon_unit not in canonical_map:
                    # First occurrence: store under the shorter representation
                    canonical_map[canon_unit] = dict(rd)
                    canonical_order.append(canon_unit)
                else:
                    # Duplicate: merge any non-null links the first entry lacks
                    existing = canonical_map[canon_unit]
                    if existing["dtcm"] is None and rd["dtcm"] is not None:
                        existing["dtcm"] = rd["dtcm"]
                    if existing["det"] is None and rd["det"] is not None:
                        existing["det"] = rd["det"]
                    # Keep the shorter unit string as the canonical display name
                    if len(raw_unit) < len(existing["unit"]):
                        existing["unit"] = raw_unit

            unit_links = []
            for canon_unit in sorted(canonical_order, key=natural_sort_key):
                merged_rd = canonical_map[canon_unit]
                unit_links.append({
                    "unit": merged_rd["unit"],
                    "dtcm": merged_rd["dtcm"],
                    "det": merged_rd["det"],
                })
                units_kept += 1

            units_in_group = [canonical_map[cu]["unit"] for cu in sorted(canonical_order, key=natural_sort_key)]

            groups_out.append({
                "_sort_key": natural_sort_key(units_in_group[0]),
                "units": units_in_group,
                "emails": rep["emails"],
                "phone": rep["phone"],
                "phoneDigits": rep["phoneDigits"],
                "portal": rep["portal"],
                "portalId": rep["portalId"],
                "portalPassword": rep["portalPassword"],
                "notes": rep["notes"],
                "unitLinks": unit_links,
            })

        # Sort groups by first unit
        groups_out.sort(key=lambda g: g["_sort_key"])

        # Strip internal sort key before output
        for g in groups_out:
            del g["_sort_key"]

        buildings_out.append({
            "name": display_name,
            "groups": groups_out,
        })

    # Sort buildings A-Z by name
    buildings_out.sort(key=lambda b: b["name"].lower())

    # -------------------------------------------------------------------------
    # Compute totals
    # -------------------------------------------------------------------------
    total_buildings = len(buildings_out)
    total_groups = sum(len(b["groups"]) for b in buildings_out)
    rows_dropped_count = len(rows_dropped)

    output = {
        "generated": date.today().isoformat(),
        "source": Path(xlsx_path).name,
        "buildings": buildings_out,
    }
    stats = {
        "rows_read": rows_read,
        "units_kept": units_kept,
        "buildings": total_buildings,
        "groups": total_groups,
        "rows_dropped": rows_dropped,
        "merged_into": merged_into,
    }
    return output, stats


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    output, stats = build_contacts(xlsx_path)

    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)

    print(f"Rows read      : {stats['rows_read']}")
    print(f"Units kept     : {stats['units_kept']}")
    print(f"Buildings      : {stats['buildings']}")
    print(f"Groups         : {stats['groups']}")
    print(f"Rows dropped   : {len(stats['rows_dropped'])}")
    for rn, reason in stats["rows_dropped"]:
        print(f"  row {rn}: {reason}")
    if stats["merged_into"]:
        print("Plural merges  :")
        for minor, canon in sorted(stats["merged_into"].items()):
            print(f"  '{minor}' merged into '{canon}'")
    print(f"Output         : {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
